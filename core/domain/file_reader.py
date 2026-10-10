"""Read CSV/XLSX files into raw sheets: a grid preview for detection and a DataFrame per
sheet once the header row is known."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

CSV_SHEET = "csv"


@dataclass
class RawSheet:
    name: str
    preview: list[list[Any]]  # first N rows, raw cell values
    n_rows: int


def _norm_header(v: Any) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"\s+", " ", str(v)).strip()


def norm_key(v: Any) -> str:
    """Case/whitespace-insensitive comparison key for column names."""
    return re.sub(r"\s+", "", _norm_header(v)).lower()


def is_excel(path: Path) -> bool:
    return path.suffix.lower() in (".xlsx", ".xlsm", ".xls")


def preview(path: Path, n: int = 30) -> list[RawSheet]:
    if is_excel(path):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows: list[list[Any]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= n:
                    break
                rows.append(list(row))
            out.append(RawSheet(ws.title, rows, ws.max_row or len(rows)))
        wb.close()
        return out
    text = _read_text(path)
    reader = csv.reader(io.StringIO(text))
    rows = []
    total = 0
    for row in reader:
        if len(rows) < n:
            rows.append(row)
        total += 1
    return [RawSheet(CSV_SHEET, rows, total)]


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def headers_at(sheet: RawSheet, header_row: int) -> list[str]:
    if header_row < 1 or header_row > len(sheet.preview):
        return []
    return [_norm_header(v) for v in sheet.preview[header_row - 1]]


def read_sheet(path: Path, sheet: str, header_row: int, raw_positions: bool = False) -> pd.DataFrame:
    """Full sheet as a DataFrame with the header at ``header_row`` (1-based).

    Columns are de-duplicated (``City``, ``City.1``). Blank header cells become
    ``col_<n>``. With ``raw_positions`` the frame keeps positional integer columns so
    an adapter can repair column shifts before naming."""
    if is_excel(path):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(min_row=header_row, values_only=True)]
        wb.close()
    else:
        rows = list(csv.reader(io.StringIO(_read_text(path))))[header_row - 1:]
    if not rows:
        return pd.DataFrame()
    header = rows[0]
    body = rows[1:]
    width = max(len(header), max((len(r) for r in body), default=0))
    body = [r + [None] * (width - len(r)) for r in body]
    names = dedupe_headers([_norm_header(h) for h in header] + [""] * (width - len(header)))
    df = pd.DataFrame(body, columns=list(range(width)) if raw_positions else names)
    if raw_positions:
        df.attrs["header_names"] = names
    # drop fully empty rows
    df = df.dropna(how="all")
    df = df[~df.apply(lambda r: all((v is None or str(v).strip() == "") for v in r), axis=1)]
    df.index = range(header_row + 1, header_row + 1 + len(df)) if len(df) else df.index
    return df


def dedupe_headers(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, n in enumerate(names):
        n = n or f"col_{i + 1}"
        if n in seen:
            seen[n] += 1
            out.append(f"{n}.{seen[n]}")
        else:
            seen[n] = 0
            out.append(n)
    return out


def infer_type(values: list[Any]) -> str:
    vals = [v for v in values if v is not None and str(v).strip() not in ("", "-")]
    if not vals:
        return "empty"
    if all(hasattr(v, "year") for v in vals):
        return "date"
    nums = 0
    dates = 0
    for v in vals:
        s = str(v).replace(",", "").strip()
        try:
            float(s)
            nums += 1
            continue
        except ValueError:
            pass
        if re.match(r"^\d{1,4}[-/ ]\w{1,3}[-/ ]\d{2,4}", s):
            dates += 1
    if nums == len(vals):
        return "number"
    if dates == len(vals):
        return "date"
    return "text"
