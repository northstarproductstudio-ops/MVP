"""Plan categories (Gaps 3–5).

Resolved once per SKU in this order, recording which source won:

1. ``mapped``   — a mapped ``plan_category`` field from any imported file (SKU master,
                   sales orders, POs or a custom file);
2. ``template`` — the formatRef template grouping;
3. ``market``   — the own-brand category column in Dcluttr market data (via the crosswalk);
4. ``none``     — uncategorised. Never an error: such SKUs are planned in general batches.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path

import pandas as pd

UNCATEGORISED = "Other"


@lru_cache(maxsize=4)
def template_layout(path: str) -> tuple[tuple[str, str, str], ...]:
    """(category, sku_code, product name) rows in template order from 'Overall DP'."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb["Overall DP"] if "Overall DP" in wb.sheetnames else wb.worksheets[0]
    out = []
    cat = UNCATEGORISED
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        a, b = (row + (None, None))[:2]
        if a in (None, "Product") and b in (None, "SKU Code"):
            continue
        if b in (None, "") and a:
            cat = str(a).strip()
            continue
        if b:
            out.append((cat, str(b).strip(), str(a or "").strip()))
    wb.close()
    return tuple(out)


def template_categories(path: Path) -> dict[str, str]:
    return {sku: cat for cat, sku, _ in template_layout(str(path))}


def plan_category_map(skus: list[str], sku_master: pd.DataFrame, template_path: Path) -> dict[str, str]:
    tmpl = template_categories(template_path) if template_path.exists() else {}
    sm = {}
    if sku_master is not None and not sku_master.empty:
        sm = {r.sku_code: r.plan_category for r in sku_master.itertuples() if r.plan_category}
    return {s: sm.get(s) or tmpl.get(s) or UNCATEGORISED for s in skus}


CATEGORY_SOURCES = ("mapped", "template", "market", "none")


def _mode_per_sku(df: pd.DataFrame, value_col: str) -> dict[str, str]:
    """Most frequent non-blank value per SKU; ties broken alphabetically (deterministic)."""
    if df is None or df.empty or value_col not in df.columns:
        return {}
    d = df[["sku_code", value_col]].dropna()
    d = d[d[value_col].astype(str).str.strip() != ""]
    if d.empty:
        return {}
    d = d.assign(v=d[value_col].astype(str).str.strip())
    counts = d.groupby(["sku_code", "v"]).size().reset_index(name="n")
    counts = counts.sort_values(["sku_code", "n", "v"], ascending=[True, False, True])
    return dict(zip(counts.drop_duplicates("sku_code")["sku_code"].astype(str),
                    counts.drop_duplicates("sku_code")["v"], strict=False))


def _mapped_from_extra(tables: dict[str, pd.DataFrame]) -> dict[str, str]:
    """plan_category values mapped in sales-order, PO or custom files (stored as extra values)."""
    ev = tables.get("extra_value")
    if ev is None or ev.empty:
        return {}
    cat = ev[ev["variable"] == "plan_category"][["dataset_id", "row_no", "value_text"]]
    if cat.empty:
        return {}
    keys = []
    for t in ("sellin_line", "po_line"):
        df = tables.get(t)
        if df is not None and not df.empty:
            keys.append(df[["dataset_id", "row_no", "sku_code"]])
    sku_ev = ev[ev["variable"] == "sku_code"][["dataset_id", "row_no", "value_text"]].rename(
        columns={"value_text": "sku_code"})
    if not sku_ev.empty:
        keys.append(sku_ev)
    if not keys:
        return {}
    k = pd.concat(keys, ignore_index=True).dropna(subset=["sku_code"]).drop_duplicates(["dataset_id", "row_no"])
    j = cat.merge(k, on=["dataset_id", "row_no"], how="inner")
    return _mode_per_sku(j, "value_text")


def resolve_plan_categories(skus: list[str], tables: dict[str, pd.DataFrame],
                            template_path: Path) -> pd.DataFrame:
    """One row per SKU: ``sku_code``, ``plan_category`` (None when uncategorised),
    ``category_source`` (mapped | template | market | none)."""
    sm = tables.get("sku_master")
    mapped: dict[str, str] = {}
    if sm is not None and not sm.empty:
        mapped.update(_mode_per_sku(sm, "plan_category"))
    for k, v in _mapped_from_extra(tables).items():
        mapped.setdefault(k, v)
    tmpl = template_categories(template_path) if template_path.exists() else {}
    md = tables.get("market_daily")
    market: dict[str, str] = {}
    if md is not None and not md.empty:
        own = md[md["is_own_brand"].astype(bool) & md["sku_code"].notna()]
        market = _mode_per_sku(own, "category")
    rows: list[tuple[str, str | None, str]] = []
    for s in sorted(set(skus)):
        if mapped.get(s):
            rows.append((s, mapped[s], "mapped"))
        elif tmpl.get(s):
            rows.append((s, tmpl[s], "template"))
        elif market.get(s):
            rows.append((s, market[s], "market"))
        else:
            rows.append((s, None, "none"))
    return pd.DataFrame(rows, columns=["sku_code", "plan_category", "category_source"])
