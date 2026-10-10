from __future__ import annotations

import json

import pandas as pd
import pytest

from core import duck
from core.domain.edits import EditError, plan_edit
from core.domain.fenced import FencedQueryError, run_fenced
from core.domain.known_queries import QUERIES, QueryParamError, run_known
from core.domain.reports import (
    current_lines,
    export_plan_xlsx,
    export_reasoning_xlsx,
    overview_facts,
    reasoning_rows,
    template_overview,
)
from core.models.edit import Edit


@pytest.fixture(scope="module")
def plan_lines(local_run):
    dp = local_run.write("dp_output")
    lines = current_lines(dp, None)
    lines["share_basis"] = "sku"
    return lines


@pytest.fixture()
def con(snapshot, plan_lines):
    c = duck.connect(None, "256MB")
    duck.register_frames(c, {k: v for k, v in snapshot.tables.items()})
    pl = plan_lines.copy()
    pl["last_month_sellin"] = 10
    pl["flags"] = pl["flags"].map(lambda f: ", ".join(f))
    duck.register_frames(c, {"plan_lines": pl})
    duck.lock_down(c)
    yield c
    c.close()


@pytest.mark.parametrize("name", sorted(QUERIES))
def test_every_known_query_runs(con, name):
    params = {"line_reasoning": {"sku": "CHCCSA35"}}.get(name, {})
    res = run_known(con, name, params)
    assert res.label
    assert isinstance(res.rows, list)


def test_known_query_param_validation(con):
    with pytest.raises(QueryParamError):
        run_known(con, "sales_history", {"grain": "year"})
    with pytest.raises(QueryParamError):
        run_known(con, "nope", {})


def test_fenced_lookup_is_labelled_and_capped(con):
    res = run_fenced(con, "select sku_code from sellin_line", row_cap=5, timeout_s=5)
    assert res.row_count == 5 and res.truncated
    assert res.label.startswith("Direct lookup")
    for bad in ("drop table sellin_line", "select * from read_parquet('x')", "select * from plan_lines"):
        with pytest.raises(FencedQueryError):
            run_fenced(con, bad, 5, 5)


def test_category_edit_and_numbers(plan_lines):
    cat = plan_lines["plan_category"].iloc[0]
    e = Edit.model_validate({"target": {"plan_category": cat}, "op": "pct", "value": 30, "note": "push"})
    p = plan_edit(plan_lines, e, 30, {}, netting_ran=False)
    sel = plan_lines[plan_lines["plan_category"] == cat]
    assert p.lines_changed == int((sel["final_dp"] > 0).sum())
    for r in p.changes.itertuples():
        assert r.dp_after == int((r.dp_before * 1.3) + 0.5)


def test_edit_unknown_names_suggest(plan_lines):
    e = Edit.model_validate({"target": {"plan_category": "Cake Decor"}, "op": "pct", "value": 10})
    with pytest.raises(EditError, match="Did you mean"):
        plan_edit(plan_lines, e, 30, {}, False)


def test_edit_schema_rules():
    with pytest.raises(ValueError):
        Edit.model_validate({"target": {}, "op": "pct", "value": 10})
    with pytest.raises(ValueError):
        Edit.model_validate({"target": {"sku_code": "X"}, "op": "pct", "value": 500})
    with pytest.raises(ValueError):
        Edit.model_validate({"target": {"sku_code": "X"}, "op": "set", "value": 10})


def test_versions_apply_in_order(plan_lines):
    lid = plan_lines["line_id"].iloc[0]
    edits = pd.DataFrame([{"version": 1, "line_id": lid, "dp_after": 100, "order_qty_after": None, "note": "a"},
                          {"version": 2, "line_id": lid, "dp_after": 120, "order_qty_after": None, "note": "b"}])
    cur = current_lines(plan_lines.drop(columns=["pipeline_dp", "user_change"]), edits)
    row = cur[cur["line_id"] == lid].iloc[0]
    assert row["final_dp"] == 120 and row["pipeline_dp"] == plan_lines.iloc[0]["final_dp"]
    assert "v1" in row["user_change"] and "v2" in row["user_change"]
    rr = reasoning_rows(cur, 2)
    assert rr.loc[rr["line_id"] == lid, "user_change"].iloc[0].startswith("v1")


def test_exports_open(local_run, plan_lines, tmp_path):
    import openpyxl

    from core.config import REPO_ROOT

    hist = local_run.ctx.artifacts["history"]
    factors = local_run.ctx.artifacts["factors"]
    meta = {"target_month": "2026-11-01", "version": 0, "run_id": "r"}
    data = export_plan_xlsx(REPO_ROOT / "templates" / "formatRef.xlsx", plan_lines, hist, factors, meta,
                            ["Bangalore", "Hyderabad", "Chennai", "Delhi", "Mumbai"])
    p = tmp_path / "plan.xlsx"
    p.write_bytes(data)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Overall DP", "Bangalore DP", "Hyderabad DP", "Chennai DP", "Delhi DP", "Mumbai DP",
                             "City shares", "Factors used for Forecast"]
    ws = wb["Overall DP"]
    vals = {r[1]: r[2] for r in ws.iter_rows(min_row=4, values_only=True) if r[1]}
    total = plan_lines.groupby("sku_code")["final_dp"].sum()
    assert vals["CHCCSA35"] == total["CHCCSA35"]
    assert vals["CHCCSA48"] == "-"  # in the template, no plan
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)
    r = export_reasoning_xlsx(reasoning_rows(plan_lines, 0), meta)
    assert len(r) > 1000
    facts = overview_facts(plan_lines, factors, 1000, "2026-11-01", 0, [])
    assert "reasoning report" in template_overview(facts)
    assert json.dumps(facts)
